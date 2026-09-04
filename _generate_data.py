# -*- coding: utf-8 -*-
"""Generate attachment/weapon seed JSON from 武器配置表-基础数据.xlsx (all 56 weapons)."""
import json
import os
import re
from collections import defaultdict

import openpyxl

SRC = 'D:/桌面/武器配置表-基础数据.xlsx'
BASE = 'D:/VS Code/bf6-gunsmith-randomizer/src/data'

SLOT_MAP = {
    '枪口': 'muzzle', '枪管': 'barrel', '下握把': 'underbarrel', '弹匣': 'magazine',
    '弹药类型': 'ammo', '瞄具': 'sight', '瞄准镜附件': 'sight_attachment',
    '顶部配件': 'top', '右侧配件': 'rail', '左侧配件': 'left_rail', '人体工学': 'ergonomics',
}

# 槽位显示顺序（与 TS 端 SLOT_ORDER 一致）
SLOT_ORDER = [
    'muzzle', 'barrel', 'underbarrel', 'magazine', 'ammo',
    'sight', 'sight_attachment', 'top', 'rail', 'left_rail', 'ergonomics',
]

ATTACHMENT_FILES = {
    'muzzle': 'muzzles.json',
    'barrel': 'barrels.json',
    'underbarrel': 'underbarrels.json',
    'magazine': 'magazines.json',
    'ammo': 'ammo-types.json',
    'sight': 'sights.json',
    'sight_attachment': 'sight-attachments.json',
    'top': 'top-attachments.json',
    'rail': 'rails.json',
    'left_rail': 'left-rails.json',
    'ergonomics': 'ergonomics.json',
}

CATEGORY_WEAPON_FILES = {
    'assault_rifle': 'assault-rifles.json',
    'carbine': 'carbines.json',
    'submachine_gun': 'submachine-guns.json',
    'light_machine_gun': 'light-machine-guns.json',
    'dmr': 'dmr.json',
    'sniper_rifle': 'sniper-rifles.json',
    'shotgun': 'shotguns.json',
}

CAT_COLS = {1: 'assault_rifle', 2: 'carbine', 3: 'submachine_gun',
            4: 'light_machine_gun', 5: 'dmr', 6: 'sniper_rifle', 7: 'shotgun'}

FIRE_MODES = {
    'assault_rifle': ['全自动', '半自动'],
    'carbine': ['全自动', '半自动'],
    'submachine_gun': ['全自动'],
    'light_machine_gun': ['全自动'],
    'dmr': ['半自动'],
    'sniper_rifle': ['栓动'],
    'shotgun': ['半自动'],
}

# 中文名 -> ASCII slug 手动映射
MANUAL_SLUG = {'隔絕者': 'gejuezhe'}

# 放大镜（瞄准镜附件）仅能与这些 1.25x~2x 瞄具搭配（排除 CQ RDS 1.25倍 / RO-S 1.25倍）
MAGNIFIER_COMPATIBLE_SIGHTS = {
    '1P87 1.5倍', '2PRO 1.25倍', '3VZR 1.75倍', 'AP2 1.75倍', 'CCO 2.00倍',
    'RO-M 1.75倍', 'ROX 1.50倍', 'R4T 2.00倍', 'SU-123 1.5倍',
}

# 必填槽位（不能为空）。其余槽位可留空（defaultEmpty=True）。
# 规则：所有武器 = 瞄具/弹匣/枪管/弹药类型；VSSM 特例 = 枪口/弹匣/瞄具/弹药类型（VSSM 无枪管）。
MANDATORY_SLOTS = {'sight', 'magazine', 'barrel', 'ammo'}
VSSM_MANDATORY_SLOTS = {'muzzle', 'magazine', 'sight', 'ammo'}

def slug(name):
    if name in MANUAL_SLUG:
        return MANUAL_SLUG[name]
    s = name.lower()
    s = re.sub(r'[^a-z0-9]+', '_', s)
    return s.strip('_')


def norm_key(s):
    """去掉标点、统一大小写，用于 sheet 名与武器名匹配（如 M60 vs M/60）。"""
    return re.sub(r'[^0-9a-z\u4e00-\u9fff]', '', s.lower())


wb = openpyxl.load_workbook(SRC, data_only=True)

# 1. 武器表：分类 -> 武器名列表（保持行序）
weapons_by_cat = defaultdict(list)
ws0 = wb['武器表']
for row in ws0.iter_rows(min_row=2):
    for c in row:
        if c.value and c.column in CAT_COLS:
            weapons_by_cat[CAT_COLS[c.column]].append(str(c.value).strip())

all_weapon_names = [n for names in weapons_by_cat.values() for n in names]
name_to_cat = {}
for cat_id, names in weapons_by_cat.items():
    for n in names:
        name_to_cat[n] = cat_id

# sheet 名 -> 武器名（处理 M60 vs M/60 这类差异）
name_lookup = {norm_key(n): n for n in all_weapon_names}

# 2. 解析所有 56 个武器 sheet
weapon_slots = {}   # 武器名 -> [(slot, [(name, cost), ...]), ...]（列顺序）
att_pool = set()    # (slot, name, cost)
for ws in wb.worksheets:
    if ws.title == '武器表':
        continue
    weapon = name_lookup.get(norm_key(ws.title))
    if weapon is None:
        raise SystemExit('无法匹配 sheet: %r' % ws.title)
    header = []
    for c in ws[1]:
        if c.value is None:
            continue
        col_name = str(c.value).split('|')[0].strip()
        header.append((c.column, col_name))
    slots = []
    for col_idx, col_name in header:
        slot = SLOT_MAP.get(col_name)
        if not slot:
            raise SystemExit('UNKNOWN SLOT %r in %s' % (col_name, weapon))
        items = []
        for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
            v = row[0].value
            if v is None or '|' not in str(v):
                continue
            name = str(v).split('|')[0].strip()
            cost = int(float(str(v).split('|')[1].strip()))
            items.append((name, cost))
            att_pool.add((slot, name, cost))
        slots.append((slot, items))
    weapon_slots[weapon] = slots

# 3. 分配配件 id（按槽位内 name, cost 排序，保证稳定）
id_map = {}        # (slot, name, cost) -> id
attachments = []   # 全部配件
by_slot = defaultdict(list)
for slot, name, cost in att_pool:
    by_slot[slot].append((name, cost))
for slot in SLOT_ORDER:
    for i, (name, cost) in enumerate(sorted(by_slot[slot], key=lambda x: (x[0], x[1])), 1):
        aid = '%s_%02d' % (slot, i)
        id_map[(slot, name, cost)] = aid
        att = {'id': aid, 'nameZh': name, 'slot': slot, 'cost': cost, 'image': ''}
        if slot == 'sight' and name in MAGNIFIER_COMPATIBLE_SIGHTS:
            att['magnifierCompatible'] = True
        attachments.append(att)

# 4. 写附件 JSON（按槽位分文件）
att_dir = os.path.join(BASE, 'attachments')
os.makedirs(att_dir, exist_ok=True)
for slot, fname in ATTACHMENT_FILES.items():
    data = [a for a in attachments if a['slot'] == slot]
    path = os.path.join(att_dir, fname)
    with open(path, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
        f.write('\n')

# 5. 生成武器对象
def build_weapon(name):
    cat = name_to_cat[name]
    mandatory = VSSM_MANDATORY_SLOTS if name == 'VSSM' else MANDATORY_SLOTS
    slots = []
    for slot, items in weapon_slots[name]:
        allowed = []
        seen = set()
        for n, c in items:
            aid = id_map.get((slot, n, c))
            if aid and aid not in seen:
                seen.add(aid)
                allowed.append(aid)
        slots.append({
            'slot': slot,
            'allowedAttachmentIds': allowed,
            'defaultEmpty': slot not in mandatory,
        })
    return {
        'id': slug(name),
        'nameZh': name,
        'category': cat,
        'maxBudget': 100,
        'fireModes': FIRE_MODES[cat],
        'slots': slots,
        'image': '',
    }

# 6. 写武器 JSON（按分类分文件，保持表内顺序）
weapon_dir = os.path.join(BASE, 'weapons')
os.makedirs(weapon_dir, exist_ok=True)
total_weapons = 0
for cat_id, fname in CATEGORY_WEAPON_FILES.items():
    data = [build_weapon(n) for n in weapons_by_cat[cat_id]]
    total_weapons += len(data)
    path = os.path.join(weapon_dir, fname)
    with open(path, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
        f.write('\n')

# 校验条件2
print('attachments:', len(attachments))
print('weapons:', total_weapons)
for n in all_weapon_names:
    slots = [s for s, _ in weapon_slots[n]]
    missing = [s for s in ('magazine', 'barrel', 'sight') if s not in slots]
    if missing:
        print('  缺槽位 %s: %s' % (n, missing))
print('magnifierCompatible 瞄具数:',
      sum(1 for a in attachments if a.get('magnifierCompatible')))
